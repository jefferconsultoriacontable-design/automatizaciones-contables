"""
Genera el informe contable y tributario mensual para gerencia, en Word.

    python generar_informe.py balance_prueba.csv --empresa "MI EMPRESA S.A.S." --nit 900123456 --mes "Marzo 2026"
    python generar_informe.py balance.xlsx --comparativo balance_febrero.xlsx

Estructura del informe
----------------------
1. Resumen ejecutivo con semaforo de hallazgos
2. Situacion financiera resumida (y variacion contra el mes anterior si se pasa --comparativo)
3. Estado de resultados y margenes
4. Indicadores financieros con su lectura
5. Situacion tributaria (saldos de impuestos por pagar y a favor)
6. Hallazgos, riesgos y recomendaciones
7. Compromisos del proximo cierre

Requiere: pip install python-docx pandas openpyxl
"""
from __future__ import annotations

import argparse
from datetime import date
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Cm

AZUL = RGBColor(0x1F, 0x3B, 0x5C)
GRIS = RGBColor(0x55, 0x5F, 0x6D)
ROJO = RGBColor(0xB3, 0x26, 0x1E)
VERDE = RGBColor(0x1E, 0x6B, 0x3A)


# --------------------------------------------------------------- datos
def cargar(ruta: Path) -> pd.DataFrame:
    df = (pd.read_excel(ruta, dtype={"codigo_puc": str}) if ruta.suffix.lower().startswith(".xls")
          else pd.read_csv(ruta, dtype={"codigo_puc": str}))
    df.columns = [c.strip().lower() for c in df.columns]
    ren = {"cuenta": "codigo_puc", "codigo": "codigo_puc", "saldo": "saldo_final",
           "nombre": "nombre_cuenta"}
    df = df.rename(columns={k: v for k, v in ren.items() if k in df.columns})
    df["saldo_final"] = pd.to_numeric(df["saldo_final"], errors="coerce").fillna(0.0)
    df["codigo_puc"] = df["codigo_puc"].astype(str)
    return df


def s(df: pd.DataFrame, *pref: str) -> float:
    return round(float(df.loc[df["codigo_puc"].str.startswith(pref), "saldo_final"].sum()), 2)


def metricas(df: pd.DataFrame) -> dict:
    m = {
        "activo": s(df, "1"),
        "activo_corriente": s(df, "11", "12", "13", "14"),
        "disponible": s(df, "11"),
        "cartera": s(df, "13"),
        "inventarios": s(df, "14"),
        "pasivo": -s(df, "2"),
        "pasivo_corriente": -s(df, "21", "22", "23", "24", "25"),
        "patrimonio": -s(df, "3"),
        "ingresos": -s(df, "41"),
        "otros_ingresos": -s(df, "42"),
        "costos": s(df, "6", "7"),
        "gastos": s(df, "5"),
        "iva_por_pagar": -s(df, "2408"),
        "retefuente_por_pagar": -s(df, "2365"),
        "reteica_por_pagar": -s(df, "2368"),
        "seguridad_social": -s(df, "2370"),
        "impuestos_a_favor": s(df, "1355"),
        "obligaciones_financieras": -s(df, "21"),
        "proveedores": -s(df, "22"),
    }
    m["utilidad_bruta"] = round(m["ingresos"] - m["costos"], 2)
    m["utilidad_operacional"] = round(m["utilidad_bruta"] - m["gastos"], 2)
    m["capital_trabajo"] = round(m["activo_corriente"] - m["pasivo_corriente"], 2)
    m["razon_corriente"] = round(m["activo_corriente"] / m["pasivo_corriente"], 2) if m["pasivo_corriente"] else 0
    m["prueba_acida"] = round((m["activo_corriente"] - m["inventarios"]) / m["pasivo_corriente"], 2) if m["pasivo_corriente"] else 0
    m["endeudamiento"] = round(m["pasivo"] / m["activo"] * 100, 2) if m["activo"] else 0
    m["margen_bruto"] = round(m["utilidad_bruta"] / m["ingresos"] * 100, 2) if m["ingresos"] else 0
    m["margen_operacional"] = round(m["utilidad_operacional"] / m["ingresos"] * 100, 2) if m["ingresos"] else 0
    m["dias_cartera"] = round(m["cartera"] * 30 / m["ingresos"], 1) if m["ingresos"] else 0
    m["dias_inventario"] = round(m["inventarios"] * 30 / m["costos"], 1) if m["costos"] else 0
    return m


def hallazgos(df: pd.DataFrame, m: dict) -> list[tuple[str, str, str]]:
    """(nivel, hallazgo, recomendacion)"""
    H = []
    if m["razon_corriente"] and m["razon_corriente"] < 1:
        H.append(("ALTO", f"Razon corriente de {m['razon_corriente']}: el pasivo corriente supera al activo corriente.",
                  "Renegociar plazos con proveedores o reprogramar deuda financiera de corto plazo."))
    if m["endeudamiento"] > 60:
        H.append(("ALTO", f"Nivel de endeudamiento del {m['endeudamiento']}%.",
                  "Evaluar capitalizacion o reduccion de pasivo financiero."))
    if m["dias_cartera"] > 60:
        H.append(("MEDIO", f"La cartera equivale a {m['dias_cartera']} dias de venta.",
                  "Depurar cartera por edades, gestionar cobro y evaluar deterioro (NIIF 9)."))
    if m["utilidad_operacional"] < 0:
        H.append(("ALTO", "Resultado operacional negativo en el periodo.",
                  "Revisar estructura de costos y gastos fijos; analizar margen por linea."))
    if s(df, "1399") == 0 and m["cartera"] > 0:
        H.append(("MEDIO", "No hay deterioro de cartera reconocido.",
                  "Estimar la perdida crediticia esperada y documentar la politica."))
    if s(df, "5160", "5260") == 0 and s(df, "15") > 0:
        H.append(("ALTO", "No se registro depreciacion del periodo teniendo PPE.",
                  "Causar la depreciacion mensual segun vida util y valor residual."))
    if m["impuestos_a_favor"] > 0:
        H.append(("MEDIO", f"Saldos a favor por impuestos de ${m['impuestos_a_favor']:,.0f}.",
                  "Verificar imputacion en la proxima declaracion o solicitar devolucion."))
    if not H:
        H.append(("BAJO", "No se identificaron hallazgos criticos en el cierre.",
                  "Mantener el control de cierre mensual y las conciliaciones al dia."))
    return H


# --------------------------------------------------------------- documento
def p(doc, texto="", size=10.5, bold=False, color=None, align=None, space_after=6):
    par = doc.add_paragraph()
    run = par.add_run(texto)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color
    if align:
        par.alignment = align
    par.paragraph_format.space_after = Pt(space_after)
    return par


def titulo(doc, texto, nivel=1):
    h = doc.add_heading(texto, level=nivel)
    for r in h.runs:
        r.font.color.rgb = AZUL
    return h


def tabla(doc, encabezados, filas, anchos=None):
    t = doc.add_table(rows=1, cols=len(encabezados))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, e in enumerate(encabezados):
        c = t.rows[0].cells[i]
        c.text = ""
        run = c.paragraphs[0].add_run(str(e))
        run.bold = True
        run.font.size = Pt(9.5)
    for fila in filas:
        celdas = t.add_row().cells
        for i, v in enumerate(fila):
            celdas[i].text = ""
            run = celdas[i].paragraphs[0].add_run(str(v))
            run.font.size = Pt(9.5)
            if i > 0:
                celdas[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if anchos:
        for fila in t.rows:
            for i, a in enumerate(anchos):
                fila.cells[i].width = Cm(a)
    return t


def money(v):
    return f"${v:,.0f}"


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("balance")
    ap.add_argument("--comparativo", default=None, help="Balance del mes anterior")
    ap.add_argument("--empresa", default="LA COMPANIA S.A.S.")
    ap.add_argument("--nit", default="900.000.000-0")
    ap.add_argument("--mes", default=date.today().strftime("%B %Y"))
    ap.add_argument("--responsable", default="Coordinador Contable y Financiero")
    ap.add_argument("--salida", default="salidas/informe_gerencia.docx")
    args = ap.parse_args()

    df = cargar(Path(args.balance))
    m = cargar(Path(args.comparativo)) if args.comparativo else None
    act, ant = metricas(df), (metricas(m) if m is not None else None)
    H = hallazgos(df, act)

    doc = Document()
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Cm(2.2)

    # -------------------------------------------------------- portada
    p(doc, args.empresa, size=18, bold=True, color=AZUL, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    p(doc, f"NIT {args.nit}", size=10, color=GRIS, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14)
    p(doc, "INFORME CONTABLE Y TRIBUTARIO MENSUAL", size=14, bold=True,
      align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    p(doc, f"Periodo: {args.mes}   |   Preparado por: {args.responsable}   |   "
           f"Fecha de emision: {date.today().strftime('%d/%m/%Y')}",
      size=9.5, color=GRIS, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=16)

    # ---------------------------------------------- 1. resumen ejecutivo
    titulo(doc, "1. Resumen ejecutivo")
    criticos = sum(1 for h in H if h[0] == "ALTO")
    p(doc, f"La compania cerro el periodo con ingresos operacionales por {money(act['ingresos'])} "
           f"y una utilidad operacional de {money(act['utilidad_operacional'])} "
           f"({act['margen_operacional']}% sobre ingresos). El activo total asciende a "
           f"{money(act['activo'])}, con un nivel de endeudamiento del {act['endeudamiento']}% "
           f"y capital de trabajo de {money(act['capital_trabajo'])}.")
    p(doc, f"Se identificaron {len(H)} hallazgos, de los cuales {criticos} son de atencion prioritaria.",
      bold=True, color=ROJO if criticos else VERDE)

    # ------------------------------------------- 2. situacion financiera
    titulo(doc, "2. Situacion financiera")
    encabezados = ["Concepto", "Actual"] + (["Mes anterior", "Variacion", "Var. %"] if ant else [])
    filas = []
    for etiqueta, llave in [("Activo total", "activo"), ("Activo corriente", "activo_corriente"),
                            ("Disponible", "disponible"), ("Cartera", "cartera"),
                            ("Inventarios", "inventarios"), ("Pasivo total", "pasivo"),
                            ("Pasivo corriente", "pasivo_corriente"), ("Patrimonio", "patrimonio"),
                            ("Capital de trabajo", "capital_trabajo")]:
        fila = [etiqueta, money(act[llave])]
        if ant:
            v = act[llave] - ant[llave]
            pct = f"{v / ant[llave] * 100:,.1f}%" if ant[llave] else "n/a"
            fila += [money(ant[llave]), money(v), pct]
        filas.append(fila)
    tabla(doc, encabezados, filas)

    # ------------------------------------------ 3. estado de resultados
    titulo(doc, "3. Estado de resultados")
    tabla(doc, ["Concepto", "Valor", "% sobre ingresos"], [
        ["Ingresos operacionales", money(act["ingresos"]), "100,0%"],
        ["(-) Costo de ventas", money(act["costos"]),
         f"{act['costos'] / act['ingresos'] * 100:,.1f}%" if act["ingresos"] else "n/a"],
        ["(=) Utilidad bruta", money(act["utilidad_bruta"]), f"{act['margen_bruto']:,.1f}%"],
        ["(-) Gastos operacionales", money(act["gastos"]),
         f"{act['gastos'] / act['ingresos'] * 100:,.1f}%" if act["ingresos"] else "n/a"],
        ["(=) Utilidad operacional", money(act["utilidad_operacional"]),
         f"{act['margen_operacional']:,.1f}%"],
        ["Otros ingresos", money(act["otros_ingresos"]), ""],
    ])

    # ------------------------------------------------ 4. indicadores
    titulo(doc, "4. Indicadores financieros")
    tabla(doc, ["Indicador", "Resultado", "Referencia", "Lectura"], [
        ["Razon corriente", act["razon_corriente"], "> 1,5",
         "Comoda" if act["razon_corriente"] >= 1.5 else "Ajustada"],
        ["Prueba acida", act["prueba_acida"], "> 1,0",
         "Comoda" if act["prueba_acida"] >= 1 else "Ajustada"],
        ["Endeudamiento", f"{act['endeudamiento']}%", "< 60%",
         "Adecuado" if act["endeudamiento"] < 60 else "Alto"],
        ["Margen bruto", f"{act['margen_bruto']}%", "Sector", ""],
        ["Margen operacional", f"{act['margen_operacional']}%", "Sector", ""],
        ["Dias de cartera", act["dias_cartera"], "< 60 dias",
         "Adecuado" if act["dias_cartera"] < 60 else "Revisar"],
        ["Dias de inventario", act["dias_inventario"], "Politica", ""],
    ])

    # --------------------------------------------- 5. situacion tributaria
    titulo(doc, "5. Situacion tributaria")
    tabla(doc, ["Concepto", "Saldo"], [
        ["IVA por pagar (neto)", money(act["iva_por_pagar"])],
        ["Retencion en la fuente por pagar", money(act["retefuente_por_pagar"])],
        ["ReteICA por pagar", money(act["reteica_por_pagar"])],
        ["Seguridad social por pagar", money(act["seguridad_social"])],
        ["Anticipos y saldos a favor", money(act["impuestos_a_favor"])],
    ])
    p(doc, "Nota: verificar que los saldos de las cuentas 2365, 2367 y 2368 coincidan con las "
           "declaraciones presentadas y con los certificados emitidos a terceros.",
      size=9, color=GRIS)

    # ------------------------------------------------- 6. hallazgos
    titulo(doc, "6. Hallazgos y recomendaciones")
    tabla(doc, ["Nivel", "Hallazgo", "Recomendacion"],
          [[n, h, r] for n, h, r in H], anchos=[1.8, 7.5, 7.5])

    # --------------------------------------------- 7. compromisos
    titulo(doc, "7. Compromisos para el proximo cierre")
    for c in ["Conciliaciones bancarias firmadas dentro de los primeros 5 dias habiles.",
              "Depuracion de cuentas por cobrar y por pagar con antiguedad superior a 90 dias.",
              "Causacion completa de nomina, prestaciones y aportes.",
              "Revision de la depreciacion y de las vidas utiles de la PPE.",
              "Cruce de las declaraciones tributarias contra los saldos contables."]:
        doc.add_paragraph(c, style="List Bullet")

    p(doc, "", space_after=12)
    p(doc, "_" * 42, size=9, color=GRIS)
    p(doc, args.responsable, size=10, bold=True, space_after=0)
    p(doc, f"{args.empresa} - {args.mes}", size=9, color=GRIS)

    salida = Path(args.salida)
    salida.parent.mkdir(parents=True, exist_ok=True)
    doc.save(salida)
    print(f"Informe generado: {salida}")
    print(f"Hallazgos: {len(H)} (prioritarios: {criticos})")


if __name__ == "__main__":
    main()
